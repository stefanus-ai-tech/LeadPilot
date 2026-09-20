import csv
from unittest.mock import Mock

import pytest

from app.batch import load_csv, parse_row, process_csv
from app.llm import LLMTimeoutError
from app.schemas import TriageResult
from app.scoring import score_lead


def write_input(path, rows, delimiter=","):
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), delimiter=delimiter)
        writer.writeheader()
        writer.writerows(rows)


def read_output(path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def runner(analysis):
    return Mock(side_effect=lambda lead: TriageResult(lead=lead, analysis=analysis,
                                                     scoring=score_lead(lead, analysis)))


def test_batch_resume_preserves_metadata_and_pending(tmp_path, lead_data, analysis):
    source = tmp_path / "input.csv"
    write_input(source, [lead_data | {"lead_id": str(i), "budget": "5000000"} for i in range(3)])
    analyze = runner(analysis)
    destination, counts = process_csv(source, tmp_path / "out", limit=1, analyze=analyze)
    assert counts == {"ok": 1, "error": 0, "pending": 2}
    assert [r["lp_status"] for r in read_output(destination / "results.csv")] == ["ok", "pending", "pending"]
    _, counts = process_csv(source, tmp_path / "out", analyze=analyze)
    assert counts == {"ok": 3, "error": 0, "pending": 0}
    assert analyze.call_count == 3
    assert [r["lead_id"] for r in read_output(destination / "results.csv")] == ["0", "1", "2"]
    assert len(read_output(destination / "errors.csv")) == 0


def test_failed_rows_continue_and_retry(tmp_path, lead_data, analysis):
    source = tmp_path / "input.csv"
    write_input(source, [lead_data, lead_data])
    good = runner(analysis)(parse_row(lead_data))
    analyze = Mock(side_effect=[LLMTimeoutError("timed out"), good])
    destination, counts = process_csv(source, tmp_path / "out", analyze=analyze)
    assert counts == {"ok": 1, "error": 1, "pending": 0}
    assert read_output(destination / "errors.csv")[0]["lp_row"] == "1"
    retry = runner(analysis)
    _, counts = process_csv(source, tmp_path / "out", analyze=retry)
    assert counts["ok"] == 2
    retry.assert_called_once()


def test_invalid_row_never_calls_model(tmp_path, lead_data, analysis):
    source = tmp_path / "input.csv"
    write_input(source, [lead_data | {"budget": "5.000.000"}, lead_data | {"budget": "0"}])
    analyze = runner(analysis)
    destination, counts = process_csv(source, tmp_path / "out", analyze=analyze)
    assert counts["error"] == 1
    assert "budget" in read_output(destination / "errors.csv")[0]["lp_error"]
    analyze.assert_called_once()


def test_semicolon_multiline_and_extra_columns(tmp_path, lead_data):
    source = tmp_path / "input.csv"
    write_input(source, [lead_data | {"message": 'Halo, "automation"\nminggu depan', "customer_id": "A1"}], ";")
    _, _, rows = load_csv(source)
    assert parse_row(rows[0]).message == 'Halo, "automation"\nminggu depan'


def test_input_change_invalidates_resume(tmp_path, lead_data, analysis):
    source = tmp_path / "input.csv"
    analyze = runner(analysis)
    write_input(source, [lead_data])
    first, _ = process_csv(source, tmp_path / "out", analyze=analyze)
    write_input(source, [lead_data | {"message": "A different requirement"}])
    second, _ = process_csv(source, tmp_path / "out", analyze=analyze)
    assert first != second
    assert analyze.call_count == 2


def test_interrupt_saves_finished_rows_and_unlocks(tmp_path, lead_data, analysis):
    source = tmp_path / "input.csv"
    write_input(source, [lead_data, lead_data])
    good = runner(analysis)(parse_row(lead_data))
    with pytest.raises(KeyboardInterrupt):
        process_csv(source, tmp_path / "out", analyze=Mock(side_effect=[good, KeyboardInterrupt()]))
    retry = runner(analysis)
    _, counts = process_csv(source, tmp_path / "out", analyze=retry)
    retry.assert_called_once()
    assert counts["ok"] == 2


def test_formula_text_is_escaped_on_export(tmp_path, lead_data, analysis):
    source = tmp_path / "input.csv"
    write_input(source, [lead_data | {"company": '=HYPERLINK("bad")'}])
    destination, _ = process_csv(source, tmp_path / "out", analyze=runner(analysis))
    assert read_output(destination / "results.csv")[0]["company"].startswith("'=")


@pytest.mark.parametrize("text", ["", "name,name,message\na,b,c", "name,message\na,hello", "name,email,message,lp_score\na,b,c,1"])
def test_bad_headers_fail_before_processing(tmp_path, text):
    source = tmp_path / "input.csv"
    source.write_text(text)
    with pytest.raises(ValueError):
        load_csv(source)


def test_wrong_column_count_is_row_error():
    with pytest.raises(ValueError):
        parse_row({"name": "A", "email": "a@example.com", "message": None})


def test_excel_export_creates_styled_sheets(tmp_path, lead_data, analysis):
    import openpyxl
    source = tmp_path / "input.csv"
    write_input(source, [
        lead_data | {"name": "Alice", "budget": "10000000", "company": '=HYPERLINK("bad")'},
        lead_data | {"name": "Bob", "budget": "invalid"},
    ])
    analyze = runner(analysis)
    destination, counts = process_csv(source, tmp_path / "out", analyze=analyze)
    assert counts["ok"] == 1
    assert counts["error"] == 1

    xlsx_path = destination / "results.xlsx"
    assert xlsx_path.exists()

    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    assert "Executive Summary" in wb.sheetnames
    assert "Lead Triage" in wb.sheetnames
    assert "Errors & Exceptions" in wb.sheetnames

    # Check Executive summary title & formulas
    ws_sum = wb["Executive Summary"]
    assert "LeadPilot" in str(ws_sum["B2"].value)
    assert ws_sum["C6"].value == "=COUNTA('Lead Triage'!A2:A3)"

    # Check Lead Triage sheet formatting
    ws_data = wb["Lead Triage"]
    assert ws_data["A1"].value == "name"
    assert ws_data["A2"].value == "Alice"
    company_col_idx = [col for col in range(1, ws_data.max_column + 1) if ws_data.cell(row=1, column=col).value == "company"][0]
    assert str(ws_data.cell(row=2, column=company_col_idx).value).startswith("'=")
    # Numeric formatting on budget
    budget_col_idx = [col for col in range(1, ws_data.max_column + 1) if ws_data.cell(row=1, column=col).value == "budget"][0]
    assert ws_data.cell(row=2, column=budget_col_idx).value == 10000000
    assert ws_data.cell(row=2, column=budget_col_idx).number_format == '"Rp" #,##0'

    # Check Errors sheet
    ws_err = wb["Errors & Exceptions"]
    assert ws_err["A2"].value == 2
    assert "budget" in ws_err["B2"].value

